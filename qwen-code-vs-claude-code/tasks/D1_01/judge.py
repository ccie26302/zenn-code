import os, sys; sys.path.insert(0, "/tasks"); sys.path.insert(0, os.path.dirname(__file__))
from jlib import ck, done; from setup import gen
from collections import defaultdict
w = sys.argv[1]; p = os.path.join(w, "report.xlsx")
st = defaultdict(int); mo = defaultdict(int)
for d, s, c, a in gen(): st[s] += a; mo[d[:7]] += a
exp_s = sorted(st.items(), key=lambda x: -x[1]); exp_m = sorted(mo.items())
if ck("report.xlsx がある", os.path.exists(p)):
    import openpyxl
    wb = openpyxl.load_workbook(p)
    ck("シート名と順序", wb.sheetnames[:2] == ["by_store", "by_month"] and len(wb.sheetnames) == 2, wb.sheetnames)
    if "by_store" in wb.sheetnames:
        r = [tuple(c.value for c in row[:2]) for row in wb["by_store"].iter_rows()]
        ck("by_store の見出しと値と順序", r[0] == ("store", "total") and [(a, int(b)) for a, b in r[1:] if a is not None] == exp_s, r[:3])
    if "by_month" in wb.sheetnames:
        ws = wb["by_month"]; r = [tuple(c.value for c in row[:2]) for row in ws.iter_rows()]
        ck("by_month の見出しと値と順序", r[0] == ("month", "total") and [(str(a), int(b)) for a, b in r[1:] if a is not None] == exp_m, r[:3])
        ck("by_month にグラフか画像がある", len(ws._charts) + len(ws._images) >= 1, (len(ws._charts), len(ws._images)))
done()
